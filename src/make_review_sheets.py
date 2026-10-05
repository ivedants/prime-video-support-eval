"""Build the two human-review workbooks from results/raw_answers.jsonl and results/scores.csv.

  review/english_review.xlsx          English-question answers with the first-pass scores, for checking.
  review/hindi_language_quality.xlsx  Hindi and Hinglish answers, for scoring language quality (0/1/2).
  review/key.json                     Maps sheet item IDs back to (model, qid, corpus) rows.

Model names are hidden in both workbooks. The English workbook includes the help-center
excerpts each answer was given, so review/ is git-ignored (article text stays local).
"""
import csv
import json
import random
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "review"
SEED = 20261002

HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
WRAP = Alignment(wrap_text=True, vertical="top")


def load():
    rows = [json.loads(l) for l in open(ROOT / "results" / "raw_answers.jsonl", encoding="utf-8")]
    scores = {(r["model"], r["qid"], r["corpus"]): r
              for r in csv.DictReader(open(ROOT / "results" / "scores.csv", encoding="utf-8"))}
    notes = {r["qid"]: r["notes"] for r in csv.DictReader(open(ROOT / "data" / "questions.csv", encoding="utf-8"))}
    chunks = {}
    for l in open(ROOT / "data" / "index" / "chunks.jsonl", encoding="utf-8"):
        d = json.loads(l)
        chunks[d["chunk_id"]] = d["text"]
    return rows, scores, notes, chunks


def style(ws, widths, input_cols, n_rows):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = WRAP
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1):
        for c in row:
            c.alignment = WRAP
            if c.column in input_cols:
                c.fill = INPUT_FILL
    ws.freeze_panes = "A2"


def english_sheet(rows, scores, notes, chunks, key):
    rng = random.Random(SEED)
    by_q = defaultdict(list)
    for r in rows:
        if r["language"] == "en":
            by_q[r["qid"]].append(r)
    wb = Workbook()
    ws = wb.active
    ws.title = "english_review"
    head = ["item", "qid", "answerable", "question", "what the articles say (note)", "excerpts the model was given",
            "answer", "first-pass correctness (0-2)", "first-pass made_up", "first-pass abstained",
            "first-pass lang_match", "first-pass reason",
            "agree? (y/n)", "your correctness", "your made_up", "your abstained", "your lang_match", "your comment"]
    ws.append(head)
    n = 0
    for qid in sorted(by_q):
        group = by_q[qid]
        rng.shuffle(group)
        for r in group:
            n += 1
            item = f"E{n:03d}"
            s = scores[(r["model"], r["qid"], r["corpus"])]
            key[item] = {"model": r["model"], "qid": r["qid"], "corpus": r["corpus"]}
            ex = "\n\n".join(f"[{i}] {chunks[c]}" for i, c in enumerate(r["retrieved_chunk_ids"], start=1))
            ws.append([item, qid, r["answerable"], r["question"], notes.get(qid, ""), ex, r["answer"],
                       int(s["correctness"]), s["made_up"], s["abstained"], s["lang_match"], s["reason"],
                       "", "", "", "", "", ""])
    style(ws, [7, 9, 10, 30, 30, 60, 70, 11, 10, 10, 10, 40, 9, 11, 10, 10, 10, 40], range(13, 19), n)
    yn = DataValidation(type="list", formula1='"y,n"', allow_blank=True)
    sc = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    ws.add_data_validation(yn)
    ws.add_data_validation(sc)
    for col in "MOPQ":
        yn.add(f"{col}2:{col}{n + 1}")
    sc.add(f"N2:N{n + 1}")
    wb.save(OUT / "english_review.xlsx")
    return n


def hindi_sheet(rows, key):
    rng = random.Random(SEED + 1)
    uniq = defaultdict(list)  # (language, pair, question, answer) -> rows
    for r in rows:
        if r["language"] in ("hi", "hinglish"):
            uniq[(r["language"], r["pair_id"], r["question"], r["answer"])].append(r)
    items = list(uniq.items())
    rng.shuffle(items)
    items.sort(key=lambda kv: (kv[0][0], kv[0][1]))  # by language, then question; random within
    wb = Workbook()
    ws = wb.active
    ws.title = "language_quality"
    ws.append(["item", "question language", "question", "answer", "language quality (0/1/2)", "comment (optional)"])
    for n, ((lang, pair, q, a), rs) in enumerate(items, start=1):
        item = f"H{n:03d}"
        key[item] = [{"model": r["model"], "qid": r["qid"], "corpus": r["corpus"]} for r in rs]
        ws.append([item, lang, q, a, "", ""])
    n = len(items)
    style(ws, [7, 10, 40, 100, 12, 40], (5, 6), n)
    sc = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    ws.add_data_validation(sc)
    sc.add(f"E2:E{n + 1}")
    wb.save(OUT / "hindi_language_quality.xlsx")
    return n, sum(len(v) for v in uniq.values())


def main():
    OUT.mkdir(exist_ok=True)
    rows, scores, notes, chunks = load()
    key = {}
    n_en = english_sheet(rows, scores, notes, chunks, key)
    n_hi, n_hi_rows = hindi_sheet(rows, key)
    json.dump(key, open(OUT / "key.json", "w"), indent=1)
    print(f"english_review.xlsx: {n_en} answers")
    print(f"hindi_language_quality.xlsx: {n_hi} distinct answers covering {n_hi_rows} result rows")


if __name__ == "__main__":
    main()
